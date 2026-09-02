import json
from io import BytesIO
import time
import uuid

import pytest
from docx import Document
from fastapi.testclient import TestClient

import db
from main import _extract_docx_text, app


db.init_db()
client = TestClient(app)


@pytest.fixture
def isolated_headers():
    headers = {"X-Local-User-Id": f"roundtrip_{uuid.uuid4().hex}"}
    yield headers
    for roadmap in client.get("/roadmaps", headers=headers).json():
        client.delete(f"/roadmaps/{roadmap['id']}", headers=headers)


def test_exported_json_file_reimport_preserves_dates_tasks_and_document(isolated_headers):
    user_id = isolated_headers["X-Local-User-Id"]
    source_id = db.save_imported_roadmap(
        f"Source {time.time_ns()}",
        [
            {
                "title": "Build the service",
                "timeframe": "Implementation",
                "granularity": "phase",
                "start_date": "2026-08-10",
                "end_date": "2026-08-14",
                "is_done": True,
                "properties": {"Priority": "High"},
                "blocks": [
                    {
                        "type": "note",
                        "data": {"text": "Keep the API stable"},
                        "position": 3,
                    }
                ],
            },
            {
                "title": "Verify the release",
                "timeframe": "Validation",
                "granularity": "milestone",
                "start_date": "2026-08-15",
                "end_date": "2026-08-16",
                "is_done": False,
                "properties": {},
                "blocks": [],
            },
        ],
        user_id,
        raw_text="Generated source roadmap",
        document=[
            {
                "type": "heading",
                "text": "Implementation",
                "level": 1,
                "section_path": ["Implementation"],
            },
            {
                "type": "table",
                "columns": ["Task", "Owner"],
                "rows": [["Build the service", "Platform"]],
                "section_path": ["Implementation"],
            },
        ],
    )
    exported = client.get(f"/roadmaps/{source_id}/export", headers=isolated_headers)
    assert exported.status_code == 200, exported.text
    source_export = exported.json()

    imported = client.post(
        "/roadmaps/import-file",
        data={"name": f"Reimported {time.time_ns()}"},
        files={
            "file": (
                "roadmap.json",
                json.dumps(source_export).encode("utf-8"),
                "application/json",
            )
        },
        headers=isolated_headers,
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["document_elements"] == len(source_export["document"])

    round_tripped = client.get(
        f"/roadmaps/{imported.json()['id']}/export",
        headers=isolated_headers,
    )
    assert round_tripped.status_code == 200, round_tripped.text

    def portable_task(task):
        return {
            key: task[key]
            for key in (
                "title",
                "timeframe",
                "granularity",
                "start_date",
                "end_date",
                "is_done",
                "properties",
            )
        } | {
            "blocks": [
                {
                    "type": block["type"],
                    "data": block["data"],
                    "position": block["position"],
                }
                for block in task["blocks"]
            ]
        }

    assert [portable_task(task) for task in round_tripped.json()["tasks"]] == [
        portable_task(task) for task in source_export["tasks"]
    ]
    assert round_tripped.json()["document"] == source_export["document"]


def test_json_document_import_enforces_the_frontend_render_contract(isolated_headers):
    payload = {
        "tasks": [{"title": "Keep the roadmap", "timeframe": "General"}],
        "document": [
            {"type": "paragraph", "text": {"unsafe": "React child"}},
            {
                "type": "paragraph",
                "text": 123,
                "section_path": ["Safe", {"unsafe": True}, 2],
            },
            {"type": "list", "items": {"unsafe": "not an array"}},
            {"type": "list", "items": ["First", {"unsafe": True}, 7, True, None]},
            {"type": "callout", "title": ["unsafe"], "text": "Readable detail"},
            {
                "type": "heading",
                "text": "Validated section",
                "level": 99,
                "page": "4",
            },
            {
                "type": "table",
                "headers": ["Task", {"unsafe": True}, 3],
                "rows": [
                    ["Ship", {"unsafe": True}],
                    {"unsafe": "row"},
                    "Single cell",
                    [None],
                ],
            },
            {"type": "unknown", "text": ["unsafe"]},
            {"type": "code", "text": False},
        ],
    }
    imported = client.post(
        "/roadmaps/import-file",
        data={"name": f"Sanitized document {time.time_ns()}"},
        files={
            "file": (
                "unsafe-document.json",
                json.dumps(payload).encode("utf-8"),
                "application/json",
            )
        },
        headers=isolated_headers,
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["document_elements"] == 6

    response = client.get(
        f"/roadmaps/{imported.json()['id']}/document",
        headers=isolated_headers,
    )
    assert response.status_code == 200, response.text
    elements = response.json()["elements"]
    assert elements == [
        {
            "type": "paragraph",
            "section_path": ["Safe", "2"],
            "text": "123",
        },
        {"type": "list", "items": ["First", "7", "true"]},
        {"type": "callout", "text": "Readable detail"},
        {"type": "heading", "page": 4, "text": "Validated section", "level": 6},
        {
            "type": "table",
            "columns": ["Task", "", "3"],
            "rows": [["Ship", ""], ["Single cell"]],
        },
        {"type": "code", "text": "false"},
    ]

    for element in elements:
        assert element["type"] in {"heading", "paragraph", "list", "callout", "code", "table"}
        for field in ("text", "title"):
            if field in element:
                assert isinstance(element[field], str)
        if "items" in element:
            assert isinstance(element["items"], list)
            assert all(isinstance(item, str) for item in element["items"])
        if element["type"] == "table":
            assert all(isinstance(column, str) for column in element["columns"])
            assert all(
                isinstance(row, list) and all(isinstance(cell, str) for cell in row)
                for row in element["rows"]
            )


@pytest.mark.parametrize(
    "payload",
    [
        {},
        {"unrelated": "value"},
        "not a roadmap",
    ],
)
def test_non_roadmap_json_shapes_are_rejected_as_empty_instead_of_crashing(
    isolated_headers,
    payload,
):
    imported = client.post(
        "/roadmaps/import-file",
        data={"name": f"Invalid JSON shape {time.time_ns()}"},
        files={
            "file": (
                "not-a-roadmap.json",
                json.dumps(payload).encode("utf-8"),
                "application/json",
            )
        },
        headers=isolated_headers,
    )

    assert imported.status_code == 400, imported.text
    assert imported.json()["detail"] == "No trackable tasks were found in this roadmap."


def test_docx_body_order_preserves_table_section_assignment(isolated_headers):
    document = Document()
    document.add_heading("Phase Alpha", level=1)
    first_table = document.add_table(rows=2, cols=2)
    first_table.cell(0, 0).text = "Task"
    first_table.cell(0, 1).text = "Owner"
    first_table.cell(1, 0).text = "Build API"
    first_table.cell(1, 1).text = "Platform"
    document.add_heading("Phase Beta", level=1)
    second_table = document.add_table(rows=2, cols=2)
    second_table.cell(0, 0).text = "Item"
    second_table.cell(0, 1).text = "Assignee"
    second_table.cell(1, 0).text = "Deploy service"
    second_table.cell(1, 1).text = "Operations"

    payload = BytesIO()
    document.save(payload)
    content = payload.getvalue()

    assert _extract_docx_text(content) == (
        "Phase Alpha\n"
        "Build API | Platform\n"
        "Phase Beta\n"
        "Deploy service | Operations"
    )

    imported = client.post(
        "/roadmaps/import-file",
        data={"name": f"Interleaved DOCX {time.time_ns()}"},
        files={
            "file": (
                "interleaved.docx",
                content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=isolated_headers,
    )
    assert imported.status_code == 200, imported.text

    tasks = client.get(
        f"/roadmaps/{imported.json()['id']}/tasks",
        headers=isolated_headers,
    )
    assert tasks.status_code == 200, tasks.text
    assert [
        (task["title"], task["timeframe_label"])
        for task in tasks.json()
    ] == [
        ("Build API | Platform", "Phase Alpha"),
        ("Deploy service | Operations", "Phase Beta"),
    ]
